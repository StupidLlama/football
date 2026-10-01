"""賽程表和裁判任務的存取（Tools shell）。

讀賽程：
- GoogleScheduleSource：secrets 有 [schedule] url（＋服務帳號）就讀 Google 試算表
- XlsxScheduleSource：否則讀本機 data/schedule.xlsx（這個檔案有人名，已被 .gitignore 排除）

存「裁判任務負責人」（可以替換，之後換資料庫只要再寫一個有 load / save 的類別）：
- GoogleDutyStore：寫進同一份賽程試算表的「裁判」分頁（服務帳號要有編輯權限）
- JsonDutyStore：本機 data/referee.json（Streamlit Cloud 重新部署會清掉，只適合本機）
"""
import json
from pathlib import Path
from typing import Protocol

from adapters.schedule import parse_fixtures
from domain.fixture import Duty, Fixture

from .config import ROOT

DEFAULT_SCHEDULE = ROOT / "data" / "schedule.xlsx"
DEFAULT_DUTIES = ROOT / "data" / "referee.json"
DUTY_HEADER = ["key", "日期", "時間", "比賽", "角色", "負責人"]


class ScheduleSource(Protocol):
    label: str

    def load_rows(self, worksheet: str) -> list[list]: ...


class XlsxScheduleSource:
    def __init__(self, path: Path = DEFAULT_SCHEDULE):
        self.path = Path(path)
        self.label = f"本機檔案 {self.path.name}"

    def load_rows(self, worksheet: str) -> list[list]:
        if not self.path.exists():
            raise FileNotFoundError(f"找不到賽程表 {self.path}（或在 secrets 設定 [schedule] url，見 README）")
        import openpyxl
        book = openpyxl.load_workbook(self.path, data_only=True, read_only=True)
        sheet = book[worksheet] if worksheet in book.sheetnames else book.worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        book.close()
        return rows


def _open_book(url: str, service_account: dict):
    import gspread
    return gspread.service_account_from_dict(dict(service_account)).open_by_url(url)


class GoogleScheduleSource:
    def __init__(self, url: str, service_account: dict):
        self.url, self.service_account = url, dict(service_account)
        self.label = "Google 試算表"

    def load_rows(self, worksheet: str) -> list[list]:
        book = _open_book(self.url, self.service_account)
        titles = [w.title for w in book.worksheets()]
        sheet = book.worksheet(worksheet) if worksheet in titles else book.sheet1
        return sheet.get_all_values()


def _google_settings(secrets):
    try:
        url = secrets["schedule"]["url"]
        sa = secrets["gcp_service_account"]
        return (url, sa) if url and sa else None
    except Exception:
        return None


def get_schedule_source(secrets=None) -> ScheduleSource:
    g = _google_settings(secrets) if secrets is not None else None
    return GoogleScheduleSource(*g) if g else XlsxScheduleSource()


def load_fixtures(source: ScheduleSource, worksheet: str) -> list[Fixture]:
    return parse_fixtures(source.load_rows(worksheet))


# ---------- 裁判任務負責人 ----------
class DutyStore(Protocol):
    label: str

    def load(self) -> dict[str, str]: ...

    def save(self, assignments: dict[str, str], duties: list[Duty]) -> None: ...


def duty_rows(assignments: dict[str, str], duties: list[Duty]) -> list[list[str]]:
    """存檔用的表格（人看得懂的欄位也一起寫，方便直接在試算表查看）。"""
    rows = [DUTY_HEADER]
    for d in duties:
        f = d.fixture
        rows.append([d.key, f.day.isoformat(), f"{f.start}–{f.end}", f"{f.home} vs {f.away}", d.role,
                     assignments.get(d.key, "")])
    return rows


def _from_rows(rows: list[list]) -> dict[str, str]:
    if not rows:
        return {}
    head = [str(c).strip() for c in rows[0]]
    if "key" not in head or "負責人" not in head:
        return {}
    k, w = head.index("key"), head.index("負責人")
    return {str(r[k]).strip(): str(r[w]).strip() for r in rows[1:] if len(r) > max(k, w) and r[k] and r[w]}


class JsonDutyStore:
    def __init__(self, path: Path = DEFAULT_DUTIES):
        self.path = Path(path)
        self.label = "本機 data/referee.json"

    def load(self) -> dict[str, str]:
        if not self.path.exists():
            return {}
        return {str(k): str(v) for k, v in json.loads(self.path.read_text(encoding="utf-8")).items() if v}

    def save(self, assignments: dict[str, str], duties: list[Duty]) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        keep = {k: v for k, v in assignments.items() if v}
        self.path.write_text(json.dumps(keep, ensure_ascii=False, indent=2), encoding="utf-8")


class GoogleDutyStore:
    def __init__(self, url: str, service_account: dict, worksheet: str):
        self.url, self.service_account, self.worksheet = url, dict(service_account), worksheet
        self.label = f"Google 試算表「{worksheet}」分頁"

    def _sheet(self, create: bool):
        book = _open_book(self.url, self.service_account)
        if self.worksheet in [w.title for w in book.worksheets()]:
            return book.worksheet(self.worksheet)
        return book.add_worksheet(self.worksheet, rows=100, cols=len(DUTY_HEADER)) if create else None

    def load(self) -> dict[str, str]:
        sheet = self._sheet(create=False)
        return _from_rows(sheet.get_all_values()) if sheet else {}

    def save(self, assignments: dict[str, str], duties: list[Duty]) -> None:
        sheet = self._sheet(create=True)
        sheet.clear()
        sheet.update(range_name="A1", values=duty_rows(assignments, duties))


def get_duty_store(secrets=None, worksheet: str = "裁判") -> DutyStore:
    g = _google_settings(secrets) if secrets is not None else None
    return GoogleDutyStore(*g, worksheet) if g else JsonDutyStore()
